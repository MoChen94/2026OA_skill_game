"""报表中心：聚合统计与 Excel 导出。"""
from collections import defaultdict
from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, Query
from openpyxl import Workbook
from sqlalchemy.orm import Session

from .. import models
from ..database import get_db
from ..deps import require_module, require_roles
from ..excel_utils import style_header, xlsx_response

router = APIRouter(prefix="/reports", tags=["报表中心"])

_MANAGERS = (models.ROLE_ADMIN, models.ROLE_DISPATCHER)


def _base_query(db: Session, user: models.User, days: int):
    q = db.query(models.WorkOrder)
    if user.role == models.ROLE_ENGINEER:
        q = q.filter(
            (models.WorkOrder.assignee_id == user.id) | (models.WorkOrder.creator_id == user.id)
        )
    if days:
        q = q.filter(models.WorkOrder.created_at >= datetime.now() - timedelta(days=days))
    return q


@router.get("/summary")
def summary(
    days: int = Query(default=30, ge=1, le=365),
    db: Session = Depends(get_db),
    user: models.User = Depends(require_module("reports")),
):
    orders = _base_query(db, user, days).order_by(models.WorkOrder.created_at).all()
    total = len(orders)
    completed = [o for o in orders if o.status == models.WO_COMPLETED]
    cancelled = [o for o in orders if o.status == models.WO_CANCELLED]
    effective = total - len(cancelled)
    completion_rate = round(len(completed) / effective * 100, 1) if effective > 0 else 0
    durations = [
        (o.finish_time - o.accept_time).total_seconds() / 3600
        for o in completed
        if o.finish_time and o.accept_time
    ]
    avg_hours = round(sum(durations) / len(durations), 1) if durations else 0

    # 类型 / 优先级分布
    type_counts: dict = defaultdict(int)
    pri_counts: dict = defaultdict(int)
    for o in orders:
        type_counts[o.order_type] += 1
        pri_counts[o.priority] += 1

    # 设备故障排行（仅故障维修类工单）
    device_faults: dict = defaultdict(int)
    device_names: dict = {}
    for o in orders:
        if o.order_type == "FAULT" and o.device_id:
            device_faults[o.device_id] += 1
            device_names[o.device_id] = f"{o.device.code} {o.device.name}" if o.device else f"#{o.device_id}"
    device_top = sorted(
        ({"device": device_names[did], "count": c} for did, c in device_faults.items()),
        key=lambda x: -x["count"],
    )[:10]

    # 工程师工作量
    eng_map: dict = {}
    for o in orders:
        if not o.assignee_id:
            continue
        s = eng_map.setdefault(o.assignee_id, {
            "name": o.assignee.real_name if o.assignee else "?",
            "title": o.assignee.title if o.assignee else "",
            "total": 0, "done": 0, "hours": [],
        })
        s["total"] += 1
        if o.status == models.WO_COMPLETED:
            s["done"] += 1
            if o.finish_time and o.accept_time:
                s["hours"].append((o.finish_time - o.accept_time).total_seconds() / 3600)
    engineer_ranking = sorted(
        (
            {
                "name": s["name"], "title": s["title"], "total": s["total"],
                "done": s["done"],
                "avg_hours": round(sum(s["hours"]) / len(s["hours"]), 1) if s["hours"] else 0,
            }
            for s in eng_map.values()
        ),
        key=lambda x: -x["done"],
    )

    # 每日新增 / 完成趋势
    created_by_day: dict = defaultdict(int)
    done_by_day: dict = defaultdict(int)
    for o in orders:
        created_by_day[o.created_at.strftime("%m-%d")] += 1
        if o.status == models.WO_COMPLETED and o.verify_time:
            done_by_day[o.verify_time.strftime("%m-%d")] += 1
    labels = [(datetime.now() - timedelta(days=i)).strftime("%m-%d") for i in range(days - 1, -1, -1)]

    return {
        "days": days,
        "total": total,
        "completed": len(completed),
        "completion_rate": completion_rate,
        "avg_hours": avg_hours,
        "type_distribution": [
            {"name": models.WO_TYPES.get(t, t), "value": c} for t, c in sorted(type_counts.items())
        ],
        "priority_distribution": [
            {"name": models.WO_PRIORITIES.get(p, p), "value": c} for p, c in sorted(pri_counts.items())
        ],
        "device_top": device_top,
        "engineer_ranking": engineer_ranking,
        "trend": {"labels": labels, "created": [created_by_day.get(l, 0) for l in labels],
                  "completed": [done_by_day.get(l, 0) for l in labels]},
    }


@router.get("/export")
def export_report(
    days: int = Query(default=30, ge=1, le=365),
    db: Session = Depends(get_db),
    user: models.User = Depends(require_roles(*_MANAGERS, module="reports")),
):
    orders = _base_query(db, user, days).order_by(models.WorkOrder.created_at.desc()).all()

    wb = Workbook()

    # 概览
    ws = wb.active
    ws.title = "概览"
    ws.append(["指标", "数值"])
    completed = sum(1 for o in orders if o.status == models.WO_COMPLETED)
    cancelled = sum(1 for o in orders if o.status == models.WO_CANCELLED)
    effective = len(orders) - cancelled
    rate = round(completed / effective * 100, 1) if effective else 0
    durations = [(o.finish_time - o.accept_time).total_seconds() / 3600 for o in orders
                 if o.status == models.WO_COMPLETED and o.finish_time and o.accept_time]
    ws.append(["工单总数", len(orders)])
    ws.append(["已完成", completed])
    ws.append(["完成率(%)", rate])
    ws.append(["平均处理时长(小时)", round(sum(durations) / len(durations), 1) if durations else 0])
    style_header(ws)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16

    # 设备故障排行
    ws2 = wb.create_sheet("设备故障排行")
    ws2.append(["设备", "故障工单数"])
    device_faults: dict = defaultdict(int)
    device_names: dict = {}
    for o in orders:
        if o.order_type == "FAULT" and o.device_id:
            device_faults[o.device_id] += 1
            device_names[o.device_id] = f"{o.device.code} {o.device.name}" if o.device else f"#{o.device_id}"
    for did, c in sorted(device_faults.items(), key=lambda x: -x[1]):
        ws2.append([device_names[did], c])
    style_header(ws2)
    ws2.column_dimensions["A"].width = 32

    # 工程师工作量
    ws3 = wb.create_sheet("工程师工作量")
    ws3.append(["工程师", "岗位", "总工单", "已完成", "平均时长(小时)"])
    eng_map: dict = {}
    for o in orders:
        if not o.assignee_id:
            continue
        s = eng_map.setdefault(o.assignee_id, {"name": o.assignee.real_name, "title": o.assignee.title or "",
                                               "total": 0, "done": 0, "hours": []})
        s["total"] += 1
        if o.status == models.WO_COMPLETED:
            s["done"] += 1
            if o.finish_time and o.accept_time:
                s["hours"].append((o.finish_time - o.accept_time).total_seconds() / 3600)
    for s in sorted(eng_map.values(), key=lambda x: -x["done"]):
        ws3.append([s["name"], s["title"], s["total"], s["done"],
                    round(sum(s["hours"]) / len(s["hours"]), 1) if s["hours"] else 0])
    style_header(ws3)
    for col, w in zip("ABCDE", [12, 22, 10, 10, 16]):
        ws3.column_dimensions[col].width = w

    # 工单明细
    ws4 = wb.create_sheet("工单明细")
    ws4.append(["工单号", "标题", "类型", "优先级", "设备", "状态",
                "指派工程师", "创建时间", "完成时间"])
    for o in orders:
        ws4.append([
            o.order_no, o.title, models.WO_TYPES.get(o.order_type, o.order_type),
            models.WO_PRIORITIES.get(o.priority, o.priority),
            o.device.name if o.device else "—",
            models.WO_STATUS_LABELS.get(o.status, o.status),
            o.assignee.real_name if o.assignee else "—",
            o.created_at.strftime("%Y-%m-%d %H:%M"),
            o.finish_time.strftime("%Y-%m-%d %H:%M") if o.finish_time else "—",
        ])
    style_header(ws4)
    for col, w in zip("ABCDEFGHI", [16, 34, 10, 8, 22, 10, 12, 16, 16]):
        ws4.column_dimensions[col].width = w

    name = f"运维报表_近{days}天_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return xlsx_response(wb, name)
